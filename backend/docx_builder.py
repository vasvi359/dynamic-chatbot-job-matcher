from docx import Document
from docx.shared import Pt


def build_ats_docx(data: dict, output_path: str):
    doc = Document()

    title = doc.add_heading(data.get("name", "Resume"), level=1)

    contact = doc.add_paragraph(data.get("contact", ""))
    contact.runs[0].font.size = Pt(10)

    if data.get("summary"):
        doc.add_heading("Summary", level=2)
        doc.add_paragraph(data["summary"])

    if data.get("skills"):
        doc.add_heading("Skills", level=2)
        doc.add_paragraph(", ".join(data["skills"]))

    if data.get("experience"):
        doc.add_heading("Experience", level=2)
        for job in data["experience"]:
            line = f"{job.get('title', '')} — {job.get('company', '')}"
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            doc.add_paragraph(job.get("dates", ""))
            for bullet in job.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if data.get("education"):
        doc.add_heading("Education", level=2)
        for edu in data["education"]:
            line = f"{edu.get('degree', '')} — {edu.get('institution', '')}"
            doc.add_paragraph(line)
            doc.add_paragraph(edu.get("dates", ""))

    doc.save(output_path)
    return output_path